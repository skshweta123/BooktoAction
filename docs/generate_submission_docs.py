from pathlib import Path
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


DOCS_DIR = Path(__file__).resolve().parent


def markdown_to_text(md_content: str) -> str:
    lines = []
    for raw_line in md_content.splitlines():
        line = raw_line.strip()
        if line.startswith("#"):
            line = line.lstrip("#").strip()
        lines.append(line)
    return "\n".join(lines).strip() + "\n"


def markdown_to_docx(md_content: str, output_path: Path) -> None:
    doc = Document()
    title_set = False

    for raw_line in md_content.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            doc.add_paragraph("")
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph(stripped[2:].strip())
            p.style = doc.styles["Title"]
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            title_set = True
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph(stripped[3:].strip())
            p.style = doc.styles["Heading 1"]
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph(stripped[4:].strip())
            p.style = doc.styles["Heading 2"]
            continue
        if stripped.startswith("#### "):
            p = doc.add_paragraph(stripped[5:].strip())
            p.style = doc.styles["Heading 3"]
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(stripped[2:].strip())
            p.style = doc.styles["List Bullet"]
            continue

        if stripped[0].isdigit() and ". " in stripped[:5]:
            p = doc.add_paragraph(stripped.split(". ", 1)[1].strip())
            p.style = doc.styles["List Number"]
            continue

        if stripped.startswith("```") or stripped == "---":
            continue

        doc.add_paragraph(stripped)

    if not title_set:
        doc.paragraphs[0].style = doc.styles["Title"]
        doc.paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.save(output_path)


def convert(md_name: str, txt_name: str, docx_name: str) -> None:
    md_path = DOCS_DIR / md_name
    text = md_path.read_text(encoding="utf-8")

    txt_path = DOCS_DIR / txt_name
    txt_path.write_text(markdown_to_text(text), encoding="utf-8")

    docx_path = DOCS_DIR / docx_name
    markdown_to_docx(text, docx_path)


if __name__ == "__main__":
    convert(
        "BOOK_TO_ACTION_IEEE_STYLE_PAPER.md",
        "BOOK_TO_ACTION_IEEE_STYLE_PAPER.txt",
        "BOOK_TO_ACTION_IEEE_STYLE_PAPER.docx",
    )
    convert(
        "BOOK_TO_ACTION_PROJECT_REPORT.md",
        "BOOK_TO_ACTION_PROJECT_REPORT.txt",
        "BOOK_TO_ACTION_PROJECT_REPORT.docx",
    )
    print("Generated TXT and DOCX files successfully.")
